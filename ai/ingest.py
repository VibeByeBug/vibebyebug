"""발표자료(PDF/PPTX) -> 슬라이드 단위 청크 JSONL.

청크 단위를 '슬라이드'로 고정한 이유:
Ready-Q가 발표자에게 내놓는 산출물이 '키워드 + 슬라이드 출처'다.
인용 단위가 곧 슬라이드이므로, 그보다 잘게 쪼개면 검색은 나아질지 몰라도
"몇 번 슬라이드" 를 되돌려줄 수 없게 된다. 검색 단위 = 인용 단위.

사용:
    python ingest.py <파일.pdf|파일.pptx> [-o data/chunks.jsonl]

종료 코드: 0 정상 / 1 파일 없음 / 3 추출 실패(텍스트 0건)
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE_TYPE


class NoTextError(RuntimeError):
    """추출 가능한 텍스트가 전혀 없는 자료. 조용히 0건으로 넘기면 안 된다."""


def _clean(text: str) -> str:
    """줄바꿈/공백 정리. 슬라이드 텍스트는 레이아웃 때문에 공백이 지저분하다."""
    text = text.replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def from_pdf(path: Path) -> list[dict]:
    import fitz  # PyMuPDF

    doc = fitz.open(path)
    chunks = []
    for i, page in enumerate(doc, start=1):
        chunks.append({"page": i, "text": _clean(page.get_text()), "notes": ""})
    doc.close()

    if not any(c["text"] for c in chunks):
        raise NoTextError(
            f"{path.name}: {len(chunks)}쪽에서 텍스트를 한 글자도 못 찾았습니다. "
            "스캔본이거나 글자가 이미지로 깔린 PDF입니다."
        )
    return [c for c in chunks if c["text"]]


def _walk(shapes):
    """도형을 재귀 순회. 그룹 안에 텍스트가 든 자료가 흔해서 반드시 파고들어야 한다."""
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)


def from_pptx(path: Path) -> list[dict]:
    from pptx import Presentation

    prs = Presentation(path)
    chunks = []
    n_picture = 0

    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in _walk(slide.shapes):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                n_picture += 1
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            # 표는 발표자료에서 수치 근거가 가장 많이 사는 곳이라 반드시 긁는다
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))

        # 발표자 노트: 청중에겐 안 보이지만 백데이터가 여기 있는 경우가 많다
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        chunks.append({"page": i, "text": _clean("\n".join(parts)), "notes": _clean(notes)})

    kept = [c for c in chunks if c["text"] or c["notes"]]
    if not kept:
        raise NoTextError(
            f"{path.name}: 슬라이드 {len(chunks)}장에서 텍스트를 한 글자도 못 찾았습니다 "
            f"(그림 도형 {n_picture}개).\n"
            "  디자인 툴에서 이미지로 내보낸 자료입니다. 텍스트 추출로는 검색이 불가능하므로,\n"
            "  슬라이드를 이미지로 렌더해 멀티모달 캡션 경로를 써야 합니다."
        )
    return kept


# ── 띄어쓰기 복원 ─────────────────────────────────────────────────────
# PowerPoint 에서 PDF 로 저장한 한글 자료 중에 공백이 통째로 빠지는 것이 있다.
#   "구독500 명으로운영BEP,"  "1회용응시링크와지원자별리포트,"
# 글자 위치를 재봐도 단어 사이 간격이 0 이거나 겹쳐서 위치로는 되살릴 수 없었다.
# 형태소 분석기(kiwipiepy, 키워드 추출에 이미 쓰는 것)로 띄어쓰기를 다시 넣는다.
#   "구독 500명으로 운영 BEP,"  "1회용 응시 링크와 지원자별 리포트,"
# 19장에 0.15초. 정상 자료에 돌리면 멀쩡한 공백을 건드리므로 공백이 모자란 쪽에만 쓴다.

# "한글 공백 한글" 이 한글 글자 수에 비해 얼마나 있나.
# 전체 공백 수로 재면 숫자, 영어, 쉼표 뒤 공백이 섞여서 공백 빠진 자료도 0.10 이 나왔다.
#   공백 빠진 PDF 0.0  /  정상 자료(톰과젤리) 0.25
MIN_SPACE_RATIO = 0.05
_kiwi = None


def _hangul(text: str) -> int:
    return len(re.findall(r"[가-힣]", text))


def needs_spacing(text: str) -> bool:
    h = _hangul(text)
    gaps = len(re.findall(r"[가-힣] (?=[가-힣])", text))
    return h >= 30 and gaps / h < MIN_SPACE_RATIO


def restore_spacing(text: str) -> str:
    global _kiwi
    if _kiwi is None:
        from kiwipiepy import Kiwi
        _kiwi = Kiwi()
    spaced = "\n".join(_kiwi.space(l, reset_whitespace=True) if _hangul(l) else l
                       for l in text.split("\n"))
    # 분석기가 숫자와 단위를 떼어놓는다 ("10 곳", "95 %"). 근거 문장에서 제일 눈에 띄는 곳이라 되붙인다.
    return _NUM_UNIT.sub(r"\1\2", spaced)


_NUM_UNIT = re.compile(r"(\d)\s+(%|곳|명|인|개|건|원|억|만|년|월|일|배|회|장|차|위|종|시간|분|초|점|퍼센트)")


# 슬라이드 글상자 폭 때문에 줄이 바뀐 자리를 되붙인다.
# 안 붙이면 근거로 "누적 응시자 1,000명," 같은 반쪽 줄이 뜬다.
_BULLET = re.compile(r"^\s*([-•▪◦●○■□※➢❖✓]|\d+[.)]|[①-⑳])")


def join_wrapped(text: str) -> str:
    out: list[str] = []
    for line in text.split("\n"):
        if out and out[-1].rstrip().endswith(",") and line.strip() and not _BULLET.match(line):
            out[-1] = out[-1].rstrip() + " " + line.strip()
        else:
            out.append(line)
    return "\n".join(out)


def ingest(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        raw = from_pdf(path)
    elif suffix == ".pptx":
        raw = from_pptx(path)
    else:
        raise ValueError(f"지원하지 않는 형식: {suffix} (.pdf 또는 .pptx만)")

    # 공백 판정은 자료 전체로 한다. 슬라이드마다 하면 짧은 장에서 흔들린다.
    spacing = needs_spacing("\n".join(c["text"] for c in raw))

    out = []
    for c in raw:
        text = restore_spacing(c["text"]) if spacing else c["text"]
        c = {**c, "text": join_wrapped(text)}
        notes = c.get("notes", "")
        body = c["text"] if not notes else f"{c['text']}\n\n[노트] {notes}"
        out.append(
            {
                "id": f"{path.stem}#p{c['page']}",
                "source": path.name,
                "page": c["page"],
                "text": body,
                "n_chars": len(body),
            }
        )
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("input", type=Path)
    ap.add_argument("-o", "--output", type=Path, default=Path("data/chunks.jsonl"))
    args = ap.parse_args()

    if not args.input.exists():
        print(f"파일 없음: {args.input}", file=sys.stderr)
        return 1

    try:
        chunks = ingest(args.input)
    except NoTextError as e:
        print(f"[추출 실패] {e}", file=sys.stderr)
        return 3

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8") as f:
        for c in chunks:
            f.write(json.dumps(c, ensure_ascii=False) + "\n")

    lens = sorted(c["n_chars"] for c in chunks)
    print(f"{args.input.name} -> {args.output}")
    print(f"  슬라이드 {len(chunks)}개")
    print(f"  글자수 min/중앙/max: {lens[0]} / {lens[len(lens) // 2]} / {lens[-1]}")
    thin = [c["page"] for c in chunks if c["n_chars"] < 30]
    if thin:
        print(f"  [주의] 30자 미만 슬라이드 {len(thin)}개: {thin[:10]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
